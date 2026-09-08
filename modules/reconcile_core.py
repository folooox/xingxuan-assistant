# -*- coding: utf-8 -*-
"""对账分析 - 数据逻辑：全量导出 → 供应商月度 账单 / 导出
- 导出（自己看）：含服务组织、金额、编码等全部信息
- 账单（给供应商）：去掉服务组织、订单总金额、余额抵扣金额、实收金额、商品编码、规格编码
- 总成本 = 成本价 × 商品数量
- 已取消订单剔除；售后单保留（列都在，人工判断）
- 未识别编码前缀汇总提示
"""
import re

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side

from core.config import load_json
from core.io_utils import smart_read
from core.suppliers import load_map, match_supplier, code_prefix

_TAIL = [f'商品内部属性{i}' for i in range(1, 11)] + \
        [f'商家备注{i}' for i in range(1, 6)] + \
        [f'自定义字段{i}' for i in range(1, 21)] + \
        ['外部订单号', '定制信息', '外部优惠金额(卡类)', '付款方信息', '外部卡优惠', '外部卡ID']

EXPORT_COLS = ['订单编号', '下单时间', '服务组织', '订单总金额', '余额抵扣金额', '实收金额',
               '发货时间', '商品名称', '商品编码', '规格编码', '商品规格', '商品数量',
               '成本价', '总成本', '预约人/收货人姓名',
               '售后类型', '售后状态', '售后件数', '退款金额', '退款时间'] + _TAIL

BILL_COLS = ['订单编号', '下单时间', '发货时间', '商品名称', '商品规格', '商品数量',
             '成本价', '总成本', '预约人/收货人姓名',
             '售后类型', '售后状态', '售后件数', '退款金额', '退款时间'] + _TAIL

THIN = Border(left=Side('thin'), right=Side('thin'), top=Side('thin'), bottom=Side('thin'))
FILL_H = PatternFill('solid', fgColor="E0E0E0")


class ReconcileProcessor:
    def __init__(self):
        self.sup_map = load_map()
        self.df = pd.DataFrame()
        self.known = {}     # 供应商名 -> 行数
        self.unknown = {}   # 前缀 -> 行数

    def load_full(self, path, log):
        self.sup_map = load_map()  # 每次加载重读映射，设置页改完即生效
        df = smart_read(path)
        total = len(df)
        if '取消类型' in df.columns:
            mask = df['取消类型'].notna() & (df['取消类型'].astype(str).str.strip() != '') \
                   & (df['取消类型'].astype(str).str.lower() != 'nan')
            df = df[~mask]
            log(f"🚫 剔除已取消: {int(mask.sum())} 条")

        # 排除不在范围内的服务组织（海亮优选/悦蓉等，见 config/exclude_orgs.json）
        excl = set(load_json("exclude_orgs.json", ["海亮优选", "悦蓉"]))
        if '服务组织' in df.columns:
            m_ex = df['服务组织'].astype(str).str.strip().isin(excl)
            if m_ex.any():
                log(f"⛔ 排除服务组织({'/'.join(excl)}): {int(m_ex.sum())} 条")
            df = df[~m_ex]

        sup_names = []
        for _, row in df.iterrows():
            sup, prefix = match_supplier(row.get('商品编码', ''), self.sup_map)
            sup_names.append(sup or '')
            if sup:
                self.known[sup] = self.known.get(sup, 0) + 1
            elif prefix:
                self.unknown[prefix] = self.unknown.get(prefix, 0) + 1
        df = df.copy()
        df['_供应商'] = sup_names
        self.df = df

        log(f"✅ 已加载 {total} 条，可对账供应商：")
        for k, v in sorted(self.known.items(), key=lambda x: -x[1]):
            log(f"   {k}: {v} 条")
        if self.unknown:
            tops = sorted(self.unknown.items(), key=lambda x: -x[1])[:15]
            log("⚠️ 未识别前缀（可能不需要，或需补 config/suppliers.json）: " +
                "、".join(f"{k}×{v}" for k, v in tops))
        return True

    def export_pair(self, supplier, bill_path, export_path, log):
        sub = self.df[self.df['_供应商'] == supplier].copy()
        if sub.empty:
            log(f"❌ 没有 {supplier} 的数据")
            return False
        sub.sort_values(by='下单时间', inplace=True)

        # 计算列与格式
        def qty(v):
            try:
                return int(float(v))
            except Exception:
                return 0

        def cost(v):
            try:
                return float(v)
            except Exception:
                return 0.0

        sub['总成本'] = [cost(a) * qty(b) for a, b in zip(sub.get('成本价', 0), sub.get('商品数量', 0))]
        sub['订单编号'] = sub['订单编号'].astype(str).str.strip().str.replace('\t', '')
        for c in ('商品编码', '规格编码'):
            if c in sub.columns:
                sub[c] = sub[c].astype(str).str.strip().str.replace('\t', '').replace('nan', '')

        self._write(sub, EXPORT_COLS, export_path, "导出")
        self._write(sub, BILL_COLS, bill_path, "账单")
        n = len(sub)
        amt = sub['总成本'].sum()
        log(f"💾 {supplier}: 导出/账单已生成，{n} 行，总成本合计 {amt:,.2f} 元")
        return True

    @staticmethod
    def _write(df, cols, path, title):
        wb = Workbook()
        ws = wb.active
        ws.title = title
        ws.append(cols)
        for _, r in df.iterrows():
            row = []
            for c in cols:
                v = r.get(c, '')
                if pd.isna(v):
                    v = ''
                row.append(v)
            ws.append(row)
        from openpyxl.utils import get_column_letter
        base_w = {'订单编号': 20, '下单时间': 19, '发货时间': 19, '商品名称': 40,
                  '商品规格': 26, '预约人/收货人姓名': 14, '服务组织': 12}
        for i, c in enumerate(cols, start=1):
            ws.column_dimensions[get_column_letter(i)].width = base_w.get(c, 10)
        for i in range(1, ws.max_row + 1):
            ws.row_dimensions[i].height = 16
        for c in ws[1]:
            c.font = Font(bold=True)
            c.fill = FILL_H
        for row in ws.iter_rows():
            for cell in row:
                cell.border = THIN
        ws.freeze_panes = 'A2'
        wb.save(path)
