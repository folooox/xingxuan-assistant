"""Export raw gateway product/store records without collapsing price ranges."""
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

HEADERS = ['商户','查询节点','节点类型','商品ID','商品名称','商品编码','创建门店',
           '最低成本价','最高成本价','成本状态','最低售价','最高售价','接口库存',
           '上下架状态','可售状态','多规格','查询时间']

def flag(value):
    return '未返回' if value is None else ('是' if value else '否')

def export_records(path, merchant, records, errors=(), scope='当前展示'):
    wb = Workbook()
    ws = wb.active
    ws.title = '商品门店成本'
    ws.append(HEADERS)
    now = datetime.now().isoformat(timespec='seconds')
    for store, row in records:
        price = row.get('goodsPrice') or {}
        lo, hi = price.get('minCostPrice'), price.get('maxCostPrice')
        status = '未返回' if lo is None and hi is None else ('部分返回' if lo is None or hi is None else ('接口返回0' if lo == hi == 0 else '已返回'))
        ws.append([merchant, store.get('vidName'), store.get('vidTypeName'), str(row.get('goodsId','')),
            row.get('title'), str(row.get('outerGoodsCode') or row.get('goodsCode') or ''), row.get('createVidName'),
            lo, hi, status, price.get('minSalePrice'), price.get('maxSalePrice'),
            (row.get('goodsStock') or {}).get('goodsStockNum'), flag(row.get('isOnline')),
            flag(row.get('isCanSell')), flag(row.get('isMultiSku')), now])
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            if isinstance(cell.value, str):
                cell.data_type = 's'
        for index in (7,8,10,11):
            row[index].number_format = '0.00'
    for cell in ws[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='244A70')
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    for i in range(1,len(HEADERS)+1):
        ws.column_dimensions[get_column_letter(i)].width = 20 if i != 5 else 48
    notes = wb.create_sheet('导出说明')
    notes.append(['导出范围',scope])
    notes.append(['行数',len(records)])
    notes.append(['读取失败门店数',len(errors)])
    notes.append(['成本口径','微盟 goodsPrice.minCostPrice / maxCostPrice 原值；多规格为商品成本区间，不是SKU明细。'])
    notes.append(['零值口径','接口返回0不代表已核实零成本；缺失值留空，不补零。'])
    notes.append(['库存口径','对应查询节点的接口库存，不合计商城和门店库存。'])
    notes.append(['数据时间','逐页实时查询，期间商品变化可能导致总数变化。'])
    if errors:
        failed = wb.create_sheet('读取失败')
        failed.append(['门店','原因'])
        for e in errors:
            failed.append([e['store'].get('vidName',''),e['error']])
        for row in failed:
            for cell in row:
                if isinstance(cell.value,str): cell.data_type='s'
    wb.save(path)
    return len(records)
