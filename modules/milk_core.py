# -*- coding: utf-8 -*-
"""牛奶团购 - 数据处理逻辑（从 milk_tool_v28 迁入，与界面分离）
改动点：
1. smart_read 移入 core.io_utils 共用
2. 供应商兜底关键词移到 config/keywords_milk.json，可改配置不改代码
3. 其余清洗/重建/导出逻辑与 v28 保持一致
"""
import os
import re
import shutil

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, PatternFill, Border, Side, Font
from openpyxl.utils.dataframe import dataframe_to_rows

from core.io_utils import smart_read
from core.config import load_json

# 供应商兜底关键词（商品名里含关键词 → 归到该供应商）
DEFAULT_KEYWORDS = {
    "一鸣": ["一鸣"],
    "好易德": ["金典", "安慕希", "伊利", "特仑苏", "蒙牛", "纯甄", "光明", "莫斯利安"],
    "翼飞": ["旺仔", "澳牧"],
    "豪阳": ["花园", "小西牛"],
}


class DataProcessor:
    def __init__(self):
        self.raw_df = pd.DataFrame()
        self.verified_df = pd.DataFrame()
        self.valid_df = pd.DataFrame()
        self.error_df = pd.DataFrame()

        self.year_term_prefix = ""
        self.date_str = ""
        self.supplier_map = {}
        self.teacher_map = {}
        self.footer_info = {"addr": "", "date": ""}
        self.extracted_images = {}
        self.temp_img_dir = os.path.join(os.getcwd(), "temp_images_cache")
        self.img_folder = None
        self.product_db = {}
        self.fallback_keywords = load_json("keywords_milk.json", DEFAULT_KEYWORDS)

    def clean_temp_dir(self):
        if os.path.exists(self.temp_img_dir):
            try:
                shutil.rmtree(self.temp_img_dir)
            except Exception:
                pass

    def load_raw_files(self, haixiao_path, waixiao_path, log_func):
        dfs = []
        for p, label in [(haixiao_path, '海小店铺'), (waixiao_path, '外小店铺')]:
            if p:
                try:
                    df = smart_read(p)
                    req_cols = ['商品名称', '商品数量']
                    if not all(c in df.columns for c in req_cols):
                        log_func(f"❌ {label}文件缺少关键列")
                        return False
                    df['Source_File'] = label
                    dfs.append(df)
                    log_func(f"✅ 读取{label}: {len(df)} 条")
                except Exception as e:
                    log_func(f"❌ 读取失败 {label}: {e}")
                    return False
        if not dfs:
            return False
        self.raw_df = pd.concat(dfs, ignore_index=True)
        return True

    def load_supplier_config(self, config_path, log_func):
        if not config_path:
            return False
        self.clean_temp_dir()
        os.makedirs(self.temp_img_dir, exist_ok=True)
        self.product_db = {}
        try:
            log_func("⏳ 解析供应商表...")
            df_config = smart_read(config_path)
            col_prod = next((c for c in df_config.columns if '品名' in str(c) or '商品' in str(c)), None)
            col_supp = next((c for c in df_config.columns if '供应商' in str(c)), None)

            row_map = {}
            if col_prod and col_supp:
                for idx, row in df_config.iterrows():
                    p = str(row[col_prod]).strip()
                    s = str(row[col_supp]).strip()
                    if p and s and p != 'nan' and s != 'nan':
                        clean_key = p.replace(" ", "").replace("\t", "").lower()
                        self.product_db[clean_key] = {'supplier': s, 'raw_name': p, 'img_path': None}
                        row_map[idx + 2] = clean_key

            tail_vals = df_config.tail(10).astype(str).values.flatten()
            for val in tail_vals:
                if 'nan' in val.lower():
                    continue
                if '收货地址' in val or '浙江省' in val:
                    self.footer_info['addr'] = val.strip()
                if '送达' in val or ('202' in val and '月' in val):
                    self.footer_info['date'] = val.strip()

            if config_path.lower().endswith(('.xlsx', '.xlsm')):
                wb = load_workbook(config_path, data_only=True)
                ws = wb.active
                if hasattr(ws, '_images'):
                    cnt = 0
                    for img in ws._images:
                        try:
                            if not hasattr(img, 'anchor'):
                                continue
                            af = getattr(img.anchor, '_from', None) or getattr(img.anchor, 'from', None)
                            at = getattr(img.anchor, '_to', None) or getattr(img.anchor, 'to', None)
                            if not af:
                                continue
                            r_start = af.row + 1
                            r_end = at.row + 1 if at else r_start

                            target = None
                            center = int((r_start + r_end) / 2)
                            if center in row_map:
                                target = row_map[center]
                            if not target:
                                target = row_map.get(r_start) or row_map.get(r_end)

                            if target and target in self.product_db:
                                fname = f"img_{cnt}.png"
                                p = os.path.join(self.temp_img_dir, fname)
                                with open(p, "wb") as f:
                                    f.write(img._data())
                                self.product_db[target]['img_path'] = p
                                cnt += 1
                        except Exception:
                            continue
                    log_func(f"🖼️ 从表格提取图片: {cnt} 张")
            return True
        except Exception as e:
            log_func(f"❌ 配置表错误: {e}")
            return False

    def load_teacher_config(self, teacher_path, log_func):
        if not teacher_path:
            self.teacher_map = {}
            return True

        self.teacher_map = {}
        try:
            log_func("⏳ 解析班主任/生活老师名单...")
            xls = pd.ExcelFile(teacher_path)
            count = 0
            for sheet_name in xls.sheet_names:
                df = pd.read_excel(xls, sheet_name=sheet_name)

                col_cls = next((c for c in df.columns if '班级' in str(c)), None)

                col_tea = None
                possible_cols = ['生活教师', '生活老师', '班主任', '姓名']
                for p_col in possible_cols:
                    found = next((c for c in df.columns if p_col in str(c)), None)
                    if found:
                        col_tea = found
                        break

                if col_cls and col_tea:
                    for _, row in df.iterrows():
                        c = str(row[col_cls]).strip()
                        if c.endswith(".0"):
                            c = c[:-2]
                        c = c.replace("班", "")

                        t = str(row[col_tea]).strip()
                        if c and t and c != 'nan' and t != 'nan':
                            self.teacher_map[c] = t
                            count += 1
            log_func(f"✅ 已加载老师信息: {count} 条")
            return True
        except Exception as e:
            log_func(f"❌ 老师表读取失败: {e}")
            return False

    # --- 阶段一：生成纯净版核对表 (全量数据) ---
    def generate_checklist(self, save_path, log_func):
        rows = []
        log_func("⚙️ 正在清洗数据(全量)...")

        for _, row in self.raw_df.iterrows():
            try:
                oid = str(row.get('订单编号', ''))
                msg = str(row.get('商品留言', '')).replace('：', ':').replace('\n', ' ').strip()
                rem = str(row.get('买家备注', '')).replace('：', ':').replace('\n', ' ').strip()

                def ext(k, t):
                    m = re.search(fr'({k})[:\s]*(.*?)(?=\s*(学校名称|学校|院舍|年级|班级|学生姓名|姓名)|$)', t)
                    return m.group(2).strip() if m else ''

                cls = ext('班级', msg)
                nam = ext('学生姓名|姓名', msg)

                r_cls = ''
                if rem and rem.lower() != 'nan':
                    m = re.search(r'(\d+\s*[A-Z]+)', rem, re.IGNORECASE)
                    r_cls = m.group(1).upper().replace(' ', '') if m else (
                        re.search(r'(\d+)\s*(?:班|$)', rem).group(1)
                        if re.search(r'(\d+)\s*(?:班|$)', rem) else '')
                f_cls = r_cls if r_cls else cls

                sch = ext('学校名称|学校', msg)
                hse = ext('院舍', msg)
                grd = ext('年级', msg)
                is_wx = False
                if re.search(r'\d+\s*[A-Za-z]', f_cls) or hse or '院舍' in msg or '院' in grd or '外' in sch:
                    is_wx = True

                d_cls = f_cls
                if not is_wx:
                    cl = f_cls.replace('班', '').replace(' ', '').upper()
                    cn = re.findall(r'\d+', cl)
                    c_int = int(cn[0]) if cn else 999
                    is_pure = cl.isdigit()

                    if is_pure:
                        if 0 < c_int // 100 < 10 and c_int < 20:
                            d_cls = f"{c_int//100}{c_int:02d}"
                        elif 100 <= c_int < 1000:
                            d_cls = str(c_int)
                        else:
                            d_cls = str(c_int)
                    else:
                        d_cls = cl

                raw_prod = str(row.get('商品名称', ''))
                prod_clean = re.sub(r'(\[.*?\]|【.*?】)', '', raw_prod).strip()
                prod_clean = prod_clean.replace('海亮外国语小学部', '').replace('海亮外国语', '').strip()

                try:
                    qty = int(float(row.get('商品数量', 0)))
                except Exception:
                    qty = 0

                rows.append({
                    '订单编号': oid,
                    '下单时间': row.get('下单时间', ''),
                    '订单状态': row.get('订单状态', ''),
                    '订单总金额': row.get('订单总金额', ''),
                    '实收金额': row.get('实收金额', ''),
                    '已付金额': row.get('已付金额', ''),
                    '余额抵扣': row.get('余额抵扣金额', ''),
                    '成本价': row.get('成本价', ''),
                    '商品单价': row.get('商品单价', ''),
                    '班级': d_cls,
                    '姓名': nam,
                    '商品留言': msg,
                    '商品全名': prod_clean,
                    '数量': qty,
                    '买家备注': rem
                })
            except Exception:
                continue

        df_check = pd.DataFrame(rows)
        if df_check.empty:
            return

        cols = ['订单编号', '下单时间', '订单状态', '订单总金额', '实收金额', '已付金额',
                '余额抵扣', '成本价', '商品单价', '班级', '姓名', '商品留言', '商品全名', '数量', '买家备注']

        for c in cols:
            if c not in df_check.columns:
                df_check[c] = ""
        df_check = df_check[cols]

        try:
            df_check.sort_values(by=['班级'], inplace=True)
        except Exception:
            pass

        wb = Workbook()
        ws = wb.active
        ws.title = "导出明细"
        ws.append(cols)
        for r in dataframe_to_rows(df_check, index=False, header=False):
            ws.append(r)

        border = Border(left=Side('thin'), right=Side('thin'), top=Side('thin'), bottom=Side('thin'))
        for row in ws.iter_rows():
            for cell in row:
                cell.border = border

        wb.save(save_path)
        log_func(f"💾 导出明细表已生成: {os.path.basename(save_path)}")

    # --- 阶段二：加载并重建数据 (含7-9年级) ---
    def load_verified_file(self, path, log_func):
        try:
            # 每次执行重读关键词配置，设置页改完即生效
            self.fallback_keywords = load_json("keywords_milk.json", DEFAULT_KEYWORDS)
            df_user = smart_read(path)
            log_func(f"✅ 加载核对表: {len(df_user)} 条数据，正在重建索引...")

            hx_names = {1: '一年级', 2: '二年级', 3: '三年级', 4: '四年级', 5: '五年级',
                        6: '六年级', 7: '七年级', 8: '八年级', 9: '九年级'}
            rebuilt_rows = []

            for _, row in df_user.iterrows():
                try:
                    d_cls = str(row.get('班级', '')).strip()
                    if d_cls.endswith(".0"):
                        d_cls = d_cls[:-2]

                    nam = str(row.get('姓名', '')).strip()
                    final_prod = str(row.get('商品全名', '')).strip()
                    qty = int(row.get('数量', 0))
                    msg = str(row.get('商品留言', ''))

                    is_wx = False
                    if re.search(r'\d+[A-Za-z]', d_cls) or '院' in msg:
                        is_wx = True

                    s_sch = 0
                    s_grp = 99
                    s_cls = 9999
                    d_grp = "未知"

                    if is_wx:
                        s_sch = 1
                        m_hse = re.search(r'(院舍)[:\s]*(.*?)(?=\s*(班级|姓名)|$)', msg)
                        h_name = m_hse.group(2).strip() if m_hse else ""
                        if not h_name:
                            md = re.search(r'\d+', d_cls)
                            if md and 1 <= int(md.group()) <= 9:
                                h_name = f"{hx_names[int(md.group())]}（外小）"
                                s_grp = int(md.group())
                            else:
                                h_name = "其他（外小）"
                        else:
                            if '森' in h_name:
                                h_name = '森林院舍'
                                s_grp = 1
                            elif '海' in h_name:
                                h_name = '海洋院舍'
                                s_grp = 2
                            elif '天' in h_name:
                                h_name = '天际院舍'
                                s_grp = 3
                        d_grp = h_name
                        mc = re.search(r'(\d+)\s*([A-Za-z]+)', d_cls)
                        if mc:
                            s_cls = int(mc.group(1)) * 1000 + ord(mc.group(2).upper()[0])
                        else:
                            s_cls = 9999
                    else:
                        s_sch = 0
                        cn = re.findall(r'\d+', d_cls)
                        c_int = int(cn[0]) if cn else 999
                        gid = 99
                        if 100 <= c_int < 1000:
                            gid = c_int // 100
                        elif 1 <= c_int <= 9:
                            gid = c_int

                        std = hx_names.get(gid, "未知")
                        suff = '小'
                        if gid >= 7 and gid != 99:
                            suff = '初'
                        d_grp = f"{std}（海{suff}）" if gid != 99 else "未知年级"
                        s_grp = gid
                        s_cls = c_int if c_int != 999 else 9999

                    prod_fp = final_prod.replace(" ", "").replace("\t", "").lower()
                    supp = "未知供应商"
                    img_path = None

                    if prod_fp in self.product_db:
                        info = self.product_db[prod_fp]
                        supp = info['supplier']
                        img_path = info['img_path']
                    else:
                        found = False
                        for db_k, info in self.product_db.items():
                            if db_k in prod_fp:
                                supp = info['supplier']
                                img_path = info['img_path']
                                found = True
                                break
                        if not found:
                            # 兜底关键词来自 config/keywords_milk.json，改配置不改代码
                            for supp_name, keys in self.fallback_keywords.items():
                                if any(str(k).lower() in prod_fp for k in keys):
                                    supp = supp_name
                                    break

                    temp_p = re.sub(r'(?<!\s)(?=\d+[mM][lL]|\d+[gG]|\*|\d+盒|\d+袋|\d+瓶)', ' ', final_prod)
                    p_short = temp_p.split(' ')[0].strip()

                    rebuilt_rows.append({
                        'Sort_School': s_sch, 'Sort_Group': s_grp, 'Sort_Class': s_cls,
                        '年级归属': d_grp, '班级': d_cls, '姓名': nam,
                        '商品简名': p_short, '商品全名': final_prod, '供应商': supp, '数量': qty,
                        'IMG': img_path, '来源': '海小店铺' if not is_wx else '外小店铺'
                    })
                except Exception:
                    continue

            self.verified_df = pd.DataFrame(rebuilt_rows)
            self.valid_df = self.verified_df[self.verified_df['Sort_Group'] != 99].copy()
            self.error_df = self.verified_df[self.verified_df['Sort_Group'] == 99].copy()

            log_func("✅ 数据重建完毕！准备导出。")
            return True
        except Exception as e:
            log_func(f"❌ 读取核对表失败: {e}")
            return False

    # --- 导出：班级明细 (仅 1-6 年级 + 海小) ---
    def export_details(self, save_path):
        df = self.valid_df[
            (self.valid_df['来源'] == '海小店铺') &
            (self.valid_df['Sort_Group'] <= 6)
        ].copy()

        if df.empty:
            return

        wb = Workbook()
        ws = wb.active
        ws.title = "班级明细"
        headers = ['年级', '班级', '姓名', '牛奶名称', '数量', '发放确认']
        fill_h = PatternFill('solid', fgColor="E0E0E0")
        fill_y = PatternFill('solid', fgColor="FFC000")
        colors = [PatternFill('solid', fgColor=c) for c in ["E6F3FF", "E6FFF2", "FFF0E6", "F2E6FF", "E6FFFF"]]
        border = Border(left=Side('thin'), right=Side('thin'), top=Side('thin'), bottom=Side('thin'))
        align = Alignment(horizontal='center', vertical='center', wrap_text=True)

        r = 1
        c_idx = 0
        df.sort_values(by=['Sort_School', 'Sort_Group', 'Sort_Class', '姓名'], inplace=True)

        for gt, grp in df.groupby('年级归属', sort=False):
            ws.append(headers)
            for c in range(1, 7):
                ws.cell(r, c).fill = fill_h
            r += 1
            g_start = r
            for ct, sub in grp.groupby('班级', sort=False):
                c_start = r
                tot = 0
                fill = colors[c_idx % 5]
                c_idx += 1

                teacher = self.teacher_map.get(str(ct), "")

                for _, row in sub.iterrows():
                    ws.append([gt, ct, row['姓名'], row['商品简名'], row['数量'], ''])
                    ws.cell(r, 2).fill = fill
                    tot += row['数量']
                    r += 1

                sum_txt = f"{teacher} 合计：{tot} 提" if teacher else f"合计：{tot} 提"
                ws.cell(r, 3, sum_txt)
                ws.merge_cells(start_row=r, end_row=r, start_column=3, end_column=5)
                for c in range(3, 6):
                    ws.cell(r, c).fill = fill_y
                ws.merge_cells(start_row=c_start, end_row=r, start_column=2, end_column=2)
                ws.cell(c_start, 2).fill = fill
                r += 1
            ws.merge_cells(start_row=g_start, end_row=r - 1, start_column=1, end_column=1)

        ws.column_dimensions['A'].width = 20
        ws.column_dimensions['B'].width = 12
        ws.column_dimensions['D'].width = 35
        for row in ws.iter_rows():
            for cell in row:
                cell.border = border
                if not cell.alignment.horizontal:
                    cell.alignment = align
        wb.save(save_path)

    # --- 导出：标签 (全量) ---
    def export_labels(self, save_path):
        df = self.valid_df.copy()
        df.sort_values(by=['Sort_School', 'Sort_Group', 'Sort_Class', '姓名'], inplace=True)
        wb = Workbook()
        ws = wb.active
        ws.title = "标签"
        ws.append(['班级', '姓名', '班级姓名', '牛奶名称', '数量'])
        for _, row in df.iterrows():
            q = int(row['数量'])
            p = row['商品简名']
            c = row['班级']
            n = row['姓名']
            if q <= 1:
                ws.append([c, n, f"{c}{n}", f"{p}*1", 1])
            else:
                for i in range(1, q + 1):
                    ws.append([c, n, f"{c}{n}", f"{p}*{q}-{i}", 1])
        wb.save(save_path)

    # --- 导出：供应商 (全量) ---
    def export_supplier(self, base_path, log_func):
        df = self.valid_df.copy()
        df = df[df['供应商'].notna() & (df['供应商'] != "未知供应商")]

        summary = df.groupby(['供应商', '商品全名', '商品简名'])['数量'].sum().reset_index()

        for supp, group in summary.groupby('供应商'):
            dname = os.path.dirname(base_path)
            fname = f"{supp}牛奶采购清单_{self.date_str}.xlsx"

            wb = Workbook()
            ws = wb.active
            ws.title = supp
            ws.append(['序号', '品名', '图片', '单位', '供应商', '海小订购数'])
            ws.column_dimensions['B'].width = 45
            ws.column_dimensions['C'].width = 12
            ws.column_dimensions['E'].width = 15
            ws.column_dimensions['F'].width = 15
            border = Border(left=Side('thin'), right=Side('thin'), top=Side('thin'), bottom=Side('thin'))
            align = Alignment(horizontal='center', vertical='center', wrap_text=True)
            for c in ws[1]:
                c.font = Font(bold=True)
                c.fill = PatternFill('solid', fgColor="E0E0E0")

            idx = 1
            r = 2
            for _, row in group.iterrows():
                ws.cell(r, 1, idx)
                ws.cell(r, 2, row['商品全名'])
                ws.cell(r, 4, "箱")
                ws.cell(r, 5, supp)
                ws.cell(r, 6, row['数量'])
                ws.row_dimensions[r].height = 60

                fp = str(row['商品全名']).replace(" ", "").replace("\t", "").lower()
                img_path = None
                if fp in self.product_db:
                    img_path = self.product_db[fp]['img_path']
                else:
                    fp_s = str(row['商品简名']).replace(" ", "").replace("\t", "").lower()
                    if fp_s in self.product_db:
                        img_path = self.product_db[fp_s]['img_path']

                if not img_path and self.img_folder:
                    try:
                        files = os.listdir(self.img_folder)
                        for k in [row['商品全名'], row['商品简名']]:
                            for f in files:
                                if str(k).lower() in f.lower():
                                    img_path = os.path.join(self.img_folder, f)
                                    break
                            if img_path:
                                break
                    except Exception:
                        pass

                if img_path and os.path.exists(img_path):
                    try:
                        img = XLImage(img_path)
                        img.height = 75
                        img.width = 75
                        ws.add_image(img, f"C{r}")
                    except Exception:
                        ws.cell(r, 3, "Err")
                else:
                    ws.cell(r, 3, "无图")
                idx += 1
                r += 1

            ws.merge_cells(start_row=r, end_row=r, start_column=1, end_column=6)
            addr = self.footer_info['addr'] if self.footer_info['addr'] else \
                "收货地址：浙江省诸暨市陶朱街道西三环路199号海亮教育园海亮小学食堂一楼南门"
            ws.cell(r, 1, addr).font = Font(bold=True)
            ws.cell(r, 1).alignment = align
            r += 1
            ws.merge_cells(start_row=r, end_row=r, start_column=1, end_column=6)
            date = self.footer_info['date'] if self.footer_info['date'] else "上午9点前送达"
            ws.cell(r, 1, date).font = Font(bold=True, color="FF0000")
            ws.cell(r, 1).alignment = align

            for row in ws.iter_rows():
                for cell in row:
                    cell.border = border
                    if not cell.alignment.horizontal:
                        cell.alignment = align

            try:
                wb.save(os.path.join(dname, fname))
                log_func(f"💾 已生成: {fname}")
            except Exception:
                log_func(f"❌ 失败: {fname}")
