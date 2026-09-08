# -*- coding: utf-8 -*-
"""订单发货 - 数据逻辑（两阶段）
阶段一：全量导出 → 发货总表（识别出的行合成一张整表，文件名=当天日期-导入文件创建时间，记录发货时间节点）
阶段二：从总表 → 按供应商拆分发货单（海亮XX订单M.D.xlsx）
- 货品名称 = 商品规格去'型号:'前缀，规格为空用商品名称；单价=成本价；合计=数量
- 已取消剔除；排除服务组织见 config/exclude_orgs.json；默认只留商家配送
- 未识别编码单独输出待核对
"""
import datetime
import os
import re

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side

from core.config import load_json
from core.io_utils import smart_read
from core.suppliers import load_map, match_supplier

MASTER_COLS = ['商品编码', '订单编号', '预约人/收货人姓名', '预约人/收货人手机号', '收货人地址',
               '商品规格', '商品数量', '成本价', '下单时间', '服务组织',
               '订单总金额', '余额抵扣金额', '实收金额', '商品名称', '规格编码']

SPLIT_HEADERS = ['日期', '订单编号', '收件人', '电话', '地址', '货品名称', '数量', '单价', '合计', '单号', '物流']

THIN = Border(left=Side('thin'), right=Side('thin'), top=Side('thin'), bottom=Side('thin'))
FILL_H = PatternFill('solid', fgColor="E0E0E0")


def master_default_name(input_path):
    """当天日期-导入文件创建时间(到分钟)：7.3-9.40；分钟为0时只写小时：6.23-13"""
    try:
        t = datetime.datetime.fromtimestamp(os.path.getctime(input_path))
    except Exception:
        t = datetime.datetime.now()
    now = datetime.datetime.now()
    ts = f"{t.hour}.{t.minute:02d}" if t.minute else f"{t.hour}"
    return f"{now.month}.{now.day}-{ts}"


def _spec_name(spec, name):
    s = str(spec or '').strip()
    if s and s.lower() != 'nan':
        s = re.sub(r'^型号[:：]\s*', '', s).strip()
        if s:
            return s
    return str(name or '').strip()


def _fmt_phone(v):
    s = str(v).strip()
    if s.endswith('.0'):
        s = s[:-2]
    return '' if s.lower() == 'nan' else s


def _clean_prefix(code):
    return str(code or '').strip().replace('\t', '').upper()


class ShipProcessor:

    # ---------- 阶段一：全量导出 → 发货总表 ----------
    def build_master(self, input_path, save_path, log, only_merchant_delivery=True):
        sup_map = load_map()
        df = smart_read(input_path)
        total = len(df)

        if '取消类型' in df.columns:
            m = df['取消类型'].notna() & (df['取消类型'].astype(str).str.strip() != '') \
                & (df['取消类型'].astype(str).str.lower() != 'nan')
            df = df[~m]
            log(f"🚫 剔除已取消: {int(m.sum())} 条")

        excl = set(load_json("exclude_orgs.json", ["海亮优选", "悦蓉"]))
        if '服务组织' in df.columns:
            m = df['服务组织'].astype(str).str.strip().isin(excl)
            if m.any():
                log(f"⛔ 排除服务组织: {int(m.sum())} 条")
            df = df[~m]

        if only_merchant_delivery and '配送方式' in df.columns:
            m = df['配送方式'].astype(str).str.contains('商家配送')
            log(f"📦 只保留商家配送: {int(m.sum())} / {len(df)} 条")
            df = df[m]

        known, unknown = [], []
        sup_cnt = {}
        for _, row in df.iterrows():
            sup, prefix = match_supplier(row.get('商品编码', ''), sup_map)
            rec = {c: row.get(c, '') for c in MASTER_COLS}
            if sup:
                rec['_sort'] = (_clean_prefix(row.get('商品编码', '')), str(row.get('订单编号', '')))
                known.append(rec)
                sup_cnt[sup] = sup_cnt.get(sup, 0) + 1
            else:
                rec['未识别前缀'] = prefix if prefix else '(空)'
                unknown.append(rec)

        known.sort(key=lambda r: r['_sort'])
        for r in known:
            r.pop('_sort', None)

        self._write(known, MASTER_COLS, save_path, "发货总表")
        log(f"💾 发货总表: {os.path.basename(save_path)}（{len(known)} 行）")
        for k, v in sorted(sup_cnt.items(), key=lambda x: -x[1]):
            log(f"   {k}: {v} 条")

        if unknown:
            up = os.path.join(os.path.dirname(save_path),
                              os.path.splitext(os.path.basename(save_path))[0] + "-未识别.xlsx")
            self._write(unknown, MASTER_COLS + ['未识别前缀'], up, "未识别")
            cnt = {}
            for it in unknown:
                cnt[it['未识别前缀']] = cnt.get(it['未识别前缀'], 0) + 1
            log(f"⚠️ 未识别 {len(unknown)} 条 → {os.path.basename(up)}，前缀: " +
                "、".join(f"{k}×{v}" for k, v in sorted(cnt.items(), key=lambda x: -x[1])[:15]))
        log(f"✅ 阶段一完成：{total} 条 → 总表 {len(known)} 行")
        return True

    # ---------- 阶段二：总表 → 各供应商发货单 ----------
    def split_master(self, master_path, out_dir, date_str, log):
        sup_map = load_map()
        df = smart_read(master_path)
        miss = [c for c in ['商品编码', '订单编号', '商品数量'] if c not in df.columns]
        if miss:
            log(f"❌ 总表缺少列: {miss}")
            return False

        year = datetime.datetime.now().year
        date_cell = f"{year}/{date_str.replace('.', '/')}"

        groups, unknown = {}, 0
        for _, row in df.iterrows():
            sup, _ = match_supplier(row.get('商品编码', ''), sup_map)
            if not sup:
                unknown += 1
                continue
            try:
                qty = int(float(row.get('商品数量', 0) or 0))
            except Exception:
                qty = 0
            groups.setdefault(sup, []).append({
                '日期': date_cell,
                '订单编号': str(row.get('订单编号', '')).strip().replace('\t', ''),
                '收件人': str(row.get('预约人/收货人姓名', '')).strip(),
                '电话': _fmt_phone(row.get('预约人/收货人手机号', '')),
                '地址': str(row.get('收货人地址', '')).strip(),
                '货品名称': _spec_name(row.get('商品规格', ''), row.get('商品名称', '')),
                '数量': qty, '单价': row.get('成本价', ''), '合计': qty,
                '单号': '', '物流': '',
            })

        for sup, items in sorted(groups.items()):
            fname = f"海亮{sup}订单{date_str}.xlsx"
            self._write(items, SPLIT_HEADERS, os.path.join(out_dir, fname), "发货单")
            log(f"💾 {fname}（{len(items)} 行）")
        if unknown:
            log(f"⚠️ 总表中有 {unknown} 行编码未识别，已跳过（请核对总表）")
        log(f"✅ 阶段二完成：{len(groups)} 个供应商发货单")
        return True

    # ---------- 写表 ----------
    @staticmethod
    def _write(items, cols, path, title):
        wb = Workbook()
        ws = wb.active
        ws.title = title
        ws.append(cols)
        for it in items:
            row = []
            for c in cols:
                v = it.get(c, '')
                if v is None or (isinstance(v, float) and pd.isna(v)):
                    v = ''
                row.append(v)
            ws.append(row)
        from openpyxl.utils import get_column_letter
        base_w = {'日期': 11, '订单编号': 20, '收件人': 10, '电话': 14, '地址': 42, '货品名称': 38,
                  '数量': 6, '单价': 9, '合计': 6, '单号': 16, '物流': 10,
                  '商品编码': 10, '预约人/收货人姓名': 12, '预约人/收货人手机号': 14, '收货人地址': 42,
                  '商品规格': 22, '商品数量': 9, '成本价': 9, '下单时间': 19, '服务组织': 13,
                  '订单总金额': 11, '余额抵扣金额': 12, '实收金额': 9, '商品名称': 40, '规格编码': 10,
                  '未识别前缀': 11}
        for i, c in enumerate(cols, start=1):
            ws.column_dimensions[get_column_letter(i)].width = base_w.get(c, 12)
        for i in range(1, ws.max_row + 1):
            ws.row_dimensions[i].height = 16
        for c in ws[1]:
            c.font = Font(bold=True)
            c.fill = FILL_H
        for row in ws.iter_rows():
            for cell in row:
                cell.border = THIN
        ws.freeze_panes = 'A2'
        wb.save(path)
