# -*- coding: utf-8 -*-
"""商品池分析 - 数据逻辑：各门店商品池导出 → 合并汇总表
输入：一个文件夹，里面是各门店的商品池列表导出（文件名=门店名，如 千城.xlsx）
输出：汇总.xlsx（与手工版列一致：选择 + 供应商门店 + 19个核心列）
"""
import os

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side

from core.io_utils import smart_read

CORE_COLS = ['在售商品id', '商品名称', '商品编码', 'SKUID', '规格编码', '规格组合',
             '商家统一价', '成本价', '实际库存', '配送方式', '商家配送模板', '商品分组',
             '实际销量', '可售状态', '上下架状态', '创建时间', '商品图片',
             '微信小程序链接', '公众号二维码链接']

THIN = Border(left=Side('thin'), right=Side('thin'), top=Side('thin'), bottom=Side('thin'))
FILL_H = PatternFill('solid', fgColor="E0E0E0")


class PoolProcessor:
    def merge(self, folder, save_path, log):
        files = [f for f in sorted(os.listdir(folder))
                 if f.lower().endswith(('.xlsx', '.xls'))
                 and not f.startswith('~$')
                 and '汇总' not in f]
        if not files:
            log("❌ 文件夹里没有找到门店导出文件")
            return False

        frames = []
        for f in files:
            try:
                df = smart_read(os.path.join(folder, f))
            except Exception as e:
                log(f"❌ 读取失败 {f}: {e}")
                continue
            store = os.path.splitext(f)[0]
            miss = [c for c in CORE_COLS if c not in df.columns]
            if miss:
                log(f"⚠️ {f} 缺少列 {miss[:3]}{'...' if len(miss) > 3 else ''}，已跳过")
                continue
            part = df[CORE_COLS].copy()
            part.insert(0, '供应商门店', store)
            part.insert(0, '选择', 1)
            frames.append(part)
            log(f"✅ {store}: {len(part)} 行")

        if not frames:
            return False
        out = pd.concat(frames, ignore_index=True)

        wb = Workbook()
        ws = wb.active
        ws.title = "汇总"
        ws.append(list(out.columns))
        for row in out.itertuples(index=False):
            ws.append(['' if pd.isna(v) else v for v in row])
        from openpyxl.utils import get_column_letter
        widths = [6, 12, 16, 45, 10, 16, 12, 22, 11, 9, 9, 12, 12, 28, 9, 9, 9, 19, 40, 40, 40]
        for i, w in enumerate(widths[:ws.max_column], start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
        for c in ws[1]:
            c.font = Font(bold=True)
            c.fill = FILL_H
        for row in ws.iter_rows():
            for cell in row:
                cell.border = THIN
        ws.freeze_panes = 'A2'
        wb.save(save_path)
        log(f"💾 汇总完成：{len(files)} 个门店 → {len(out)} 行 → {os.path.basename(save_path)}")
        return True
