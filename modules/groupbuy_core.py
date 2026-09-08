# -*- coding: utf-8 -*-
"""团购配送 - 数据处理逻辑（山姆/牛奶等微盟导出的团购订单）
流程与牛奶助手一致：
阶段一：原始导出(csv/xlsx) → 剔除已取消 → 抽取核心列 + 地址清洗 → 核对表
阶段二：人工核对后的表 → 全局排序 → 团购明细 / 团购标签（标签不含自提）
"""
import os
import re

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, PatternFill, Border, Side, Font

from core.config import load_json
from core.io_utils import smart_read

# 阶段一保留的核心列（大致位置/配送位置由商品留言解析而来）
CORE_COLS = ['订单编号', '下单时间', '订单总金额', '余额抵扣金额', '实收金额',
             '大致位置', '配送位置', '商品名称', '商品数量', '成本价',
             '预约人/收货人姓名', '预约人/收货人手机号',
             '售后类型', '售后状态', '售后件数', '退款金额']

# 清洗口径默认值，可在 设置→团购地址清洗 里改（config/groupbuy.json）
_GB_DEFAULT = {
    "store_tag_raw": "教师公寓便利店自提",
    "store_name": "教工店",
    "apart_name": "教师公寓",
    "pickup_words": ['自提', '便利店', '超市', '教工店', '门店', '商店', '店里', '到店'],
}
STORE_TAG_RAW = _GB_DEFAULT["store_tag_raw"]
STORE_NAME = _GB_DEFAULT["store_name"]
APART_NAME = _GB_DEFAULT["apart_name"]
PICKUP_WORDS = _GB_DEFAULT["pickup_words"]


def refresh_config():
    """每次执行任务前重读配置，设置页改完即生效"""
    global STORE_TAG_RAW, STORE_NAME, APART_NAME, PICKUP_WORDS
    c = load_json("groupbuy.json", _GB_DEFAULT)
    STORE_TAG_RAW = c.get("store_tag_raw", _GB_DEFAULT["store_tag_raw"])
    STORE_NAME = c.get("store_name", _GB_DEFAULT["store_name"])
    APART_NAME = c.get("apart_name", _GB_DEFAULT["apart_name"])
    PICKUP_WORDS = c.get("pickup_words", _GB_DEFAULT["pickup_words"])

CN_NUM = {'一': '1', '二': '2', '三': '3', '四': '4', '五': '5',
          '六': '6', '七': '7', '八': '8', '九': '9', '两': '2', '十': '10'}

THIN = Border(left=Side('thin'), right=Side('thin'), top=Side('thin'), bottom=Side('thin'))
CENTER = Alignment(horizontal='center', vertical='center', wrap_text=False)
FILL_H = PatternFill('solid', fgColor="E0E0E0")


# ---------------- 地址解析与清洗 ----------------

def parse_message(msg):
    """'大概位置/选择自提:教师公寓\\n具体位置/填写自提:1栋一单元' → (大致, 具体)"""
    t = str(msg or '').replace('：', ':').replace('\n', ' ')
    m = re.search(r'大概位置/选择自提:?\s*(.*?)\s*具体位置/填写自提:?\s*(.*)$', t)
    if m:
        return m.group(1).strip(), m.group(2).strip()
    m = re.search(r'大概位置/选择自提:?\s*(.*)$', t)
    if m:
        return m.group(1).strip(), ''
    return '', t.strip()


def _cn2digit(s):
    for k, v in CN_NUM.items():
        s = s.replace(k, v)
    return s


def looks_like_apartment(text):
    """判断配送位置是不是教师公寓类地址（带楼栋/单元/房号特征）"""
    t = _cn2digit(str(text)).replace('幢', '栋')
    if any(k in t for k in ['栋', '单元', '室']):
        return True
    # 类似 11-2 / 11-1-1401 的纯数字段
    if re.fullmatch(r'\s*\d{1,2}\s*[-—–]\s*\d{1,4}(\s*[-—–]\s*\d{1,5})?\s*', t):
        return True
    return False


def normalize_apartment(text):
    """统一成 '11栋2单元901' 格式，方便按数字筛选"""
    t = _cn2digit(str(text)).replace('幢', '栋')
    nums = re.findall(r'\d+', t)
    if not nums:
        return str(text).strip()
    if len(nums) >= 3:
        return f"{int(nums[0])}栋{int(nums[1])}单元{nums[2]}"
    if len(nums) == 2:
        # 第二个数是3位以上→当房号；否则当单元
        if len(nums[1]) >= 3:
            return f"{int(nums[0])}栋{nums[1]}"
        return f"{int(nums[0])}栋{int(nums[1])}单元"
    return f"{int(nums[0])}栋"


def is_pickup_text(text):
    """自提类文本：含自提/便利店/超市等字眼且不带门牌数字（'公寓便利店'算自提）"""
    t = str(text).strip()
    if not t or t.lower() == 'nan':
        return False
    if re.search(r'\d', _cn2digit(t)):
        return False
    return any(w in t for w in PICKUP_WORDS)


def clean_address(tag, loc):
    """以配送位置为基准归位：返回 (大致位置, 配送位置)"""
    tag = str(tag).strip()
    loc = str(loc).strip()
    if loc.lower() == 'nan':
        loc = ''
    if tag == STORE_TAG_RAW:
        tag = STORE_NAME

    if tag in (STORE_NAME, APART_NAME):
        # 教工店/教师公寓之间以配送位置文本为准互相纠偏
        if is_pickup_text(loc) or (not loc and tag == STORE_NAME):
            return STORE_NAME, '自提'
        if looks_like_apartment(loc):
            return APART_NAME, normalize_apartment(loc)
        if '公寓' in loc or '教师公' in loc:
            # 写了公寓但没写清楼栋的，归教师公寓、原文保留待人工补
            return APART_NAME, loc
        return tag, loc

    # 其他楼栋/学校标签：照抄，只处理明显的公寓地址误选
    if looks_like_apartment(loc):
        return APART_NAME, normalize_apartment(loc)
    return tag, loc


def addr_sort_key(tag, loc):
    """教工店最前，教师公寓其次，其余按标签排；公寓地址按 栋/单元/房号 数字排"""
    if tag == STORE_NAME:
        grp = 0
    elif tag == APART_NAME:
        grp = 1
    else:
        grp = 2
    nums = re.findall(r'\d+', _cn2digit(str(loc)))
    num_key = tuple(int(n) for n in nums[:3]) + (0,) * (3 - min(len(nums), 3))
    return (grp, str(tag), num_key, str(loc))


def strip_prefix(name):
    """去掉商品名开头的【山姆团购】这类前缀（含漏写右括号的残缺前缀）"""
    s = re.sub(r'^(\s*(\[.*?\]|【.*?】))+', '', str(name)).strip()
    # 残缺前缀：如 '【山姆团购Member's...' 缺少】
    s = re.sub(r'^【[^】]{0,10}?(团购|拼团|专享)】?', '', s).strip()
    return s


def _fmt_phone(v):
    s = str(v).strip()
    if s.endswith('.0'):
        s = s[:-2]
    return '' if s.lower() == 'nan' else s


# ---------------- 处理器 ----------------

class GroupbuyProcessor:
    def __init__(self):
        self.checked_df = pd.DataFrame()

    # --- 阶段一：原始导出 → 核对表 ---
    def generate_checklist(self, input_path, save_path, log):
        refresh_config()
        df = smart_read(input_path)
        total = len(df)

        # 剔除已取消（取消类型非空）
        if '取消类型' in df.columns:
            cancel_mask = df['取消类型'].notna() & (df['取消类型'].astype(str).str.strip() != '') \
                          & (df['取消类型'].astype(str).str.lower() != 'nan')
            df = df[~cancel_mask]
            log(f"🚫 剔除已取消订单: {int(cancel_mask.sum())} 条")

        rows = []
        for _, row in df.iterrows():
            tag_raw, loc_raw = parse_message(row.get('商品留言', ''))
            tag, loc = clean_address(tag_raw, loc_raw)
            rows.append({
                '订单编号': str(row.get('订单编号', '')).strip(),
                '下单时间': row.get('下单时间', ''),
                '订单总金额': row.get('订单总金额', ''),
                '余额抵扣金额': row.get('余额抵扣金额', ''),
                '实收金额': row.get('实收金额', ''),
                '大致位置': tag,
                '配送位置': loc,
                '商品名称': strip_prefix(row.get('商品名称', '')),
                '商品数量': row.get('商品数量', ''),
                '成本价': row.get('成本价', ''),
                '预约人/收货人姓名': str(row.get('预约人/收货人姓名', '')).strip(),
                '预约人/收货人手机号': _fmt_phone(row.get('预约人/收货人手机号', '')),
                '售后类型': row.get('售后类型', ''),
                '售后状态': row.get('售后状态', ''),
                '售后件数': row.get('售后件数', ''),
                '退款金额': row.get('退款金额', ''),
            })

        out = pd.DataFrame(rows, columns=CORE_COLS)
        n_after = len(out)
        n_as = int((out['售后类型'].notna()
                    & (out['售后类型'].astype(str).str.strip() != '')
                    & (out['售后类型'].astype(str).str.lower() != 'nan')).sum())

        self._write_sheet(out, save_path, title="核对表")
        log(f"💾 核对表已生成: {os.path.basename(save_path)}")
        log(f"   共 {total} 条 → 保留 {n_after} 条，其中售后类 {n_as} 条请手工判断")
        return True

    # --- 阶段二：核对后的表 → 明细/标签 ---
    def load_verified(self, path, log):
        refresh_config()
        df = smart_read(path)
        need = ['大致位置', '配送位置', '商品名称', '商品数量', '预约人/收货人姓名', '预约人/收货人手机号']
        miss = [c for c in need if c not in df.columns]
        if miss:
            log(f"❌ 核对表缺少列: {miss}")
            return False
        df = df.copy()
        df['_key'] = df.apply(lambda r: addr_sort_key(r['大致位置'], r['配送位置']), axis=1)
        df.sort_values(by=['_key', '预约人/收货人姓名', '商品名称'], inplace=True)
        self.checked_df = df
        log(f"✅ 已加载核对表 {len(df)} 条，排序完成")
        return True

    def export_details(self, save_path, log):
        headers = ['大致地址', '配送地址', '顾客姓名', '联系电话', '商品名称', '份数', '发放确认']
        wb = Workbook()
        ws = wb.active
        ws.title = "团购明细"
        ws.append(headers)
        for _, r in self.checked_df.iterrows():
            ws.append([r['大致位置'], r['配送位置'], r['预约人/收货人姓名'],
                       _fmt_phone(r['预约人/收货人手机号']), r['商品名称'],
                       int(float(r['商品数量'] or 0)), ''])
        self._style_sheet(ws, widths=[14, 18, 12, 15, 40, 8, 10])
        wb.save(save_path)
        log(f"💾 明细已生成: {os.path.basename(save_path)}（{ws.max_row - 1} 行）")

    def export_labels(self, save_path, log):
        headers = ['大致地址', '配送地址', '顾客姓名', '联系电话', '商品名称', '份数', '发放确认']
        wb = Workbook()
        ws = wb.active
        ws.title = "团购标签"
        ws.append(headers)
        n_skip = 0
        for _, r in self.checked_df.iterrows():
            if str(r['大致位置']).strip() == STORE_NAME:
                n_skip += 1
                continue  # 自提不打标签
            q = int(float(r['商品数量'] or 0))
            base = [r['大致位置'], r['配送位置'], r['预约人/收货人姓名'],
                    _fmt_phone(r['预约人/收货人手机号'])]
            if q <= 1:
                ws.append(base + [f"{r['商品名称']}*1", 1, ''])
            else:
                for i in range(1, q + 1):
                    ws.append(base + [f"{r['商品名称']}*{q}-{i}", 1, ''])
        self._style_sheet(ws, widths=[14, 18, 12, 15, 42, 8, 10])
        wb.save(save_path)
        log(f"💾 标签已生成: {os.path.basename(save_path)}（{ws.max_row - 1} 张，跳过自提 {n_skip} 条）")

    # --- 内部：写表+样式 ---
    def _write_sheet(self, df, save_path, title):
        wb = Workbook()
        ws = wb.active
        ws.title = title
        ws.append(list(df.columns))
        for row in df.itertuples(index=False):
            ws.append(list(row))
        widths = [18, 19, 11, 11, 9, 15, 18, 40, 9, 9, 12, 14, 9, 9, 9, 9]
        self._style_sheet(ws, widths=widths)
        wb.save(save_path)

    @staticmethod
    def _style_sheet(ws, widths=None):
        from openpyxl.utils import get_column_letter
        if widths:
            for i, w in enumerate(widths[:ws.max_column], start=1):
                ws.column_dimensions[get_column_letter(i)].width = w
        for i in range(1, ws.max_row + 1):
            ws.row_dimensions[i].height = 16  # 固定16磅
        for c in ws[1]:
            c.font = Font(bold=True)
            c.fill = FILL_H
        for row in ws.iter_rows():
            for cell in row:
                cell.border = THIN
        ws.freeze_panes = 'A2'
