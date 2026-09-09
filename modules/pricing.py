"""Purchasing price import and explicit local product bindings."""
import json
import tkinter as tk
from datetime import date, datetime
from pathlib import Path
from tkinter import ttk, filedialog, messagebox

from core.config import app_dir
from modules.pricing_core import (FIELDS, load_tables, suggest_mapping, parse_quotes,
    candidate_products, product_key, compare_quote, expiry_state, text)


def table(parent, columns, height=9):
    box = ttk.Frame(parent)
    box.pack(fill='both', expand=True, pady=4)
    tree = ttk.Treeview(box, columns=columns, show='headings', height=height)
    for col in columns:
        tree.heading(col, text=col)
        tree.column(col, width=125, minwidth=70)
    y = ttk.Scrollbar(box, command=tree.yview)
    x = ttk.Scrollbar(box, orient='horizontal', command=tree.xview)
    tree.configure(yscrollcommand=y.set, xscrollcommand=x.set)
    tree.grid(row=0, column=0, sticky='nsew')
    y.grid(row=0, column=1, sticky='ns')
    x.grid(row=1, column=0, sticky='ew')
    box.rowconfigure(0, weight=1)
    box.columnconfigure(0, weight=1)
    return tree


class ImportDialog(tk.Toplevel):
    def __init__(self, parent, path, callback):
        super().__init__(parent)
        self.title('导入采购定价表：确认工作表和列对应关系')
        self.geometry('980x750')
        self.path, self.callback = path, callback
        self.tables = load_tables(path)
        self.mapping_boxes = {}
        top = ttk.Frame(self, padding=10)
        top.pack(fill='x')
        ttk.Label(top, text='工作表').pack(side='left')
        self.sheet = ttk.Combobox(top, values=list(self.tables), state='readonly', width=35)
        self.sheet.current(0)
        self.sheet.pack(side='left', padx=5)
        ttk.Label(top, text='表头在第几行').pack(side='left')
        self.header = ttk.Spinbox(top, from_=1, to=100, width=5)
        self.header.set('1')
        self.header.pack(side='left', padx=5)
        ttk.Button(top, text='刷新预览', command=self.preview).pack(side='left')
        self.sheet.bind('<<ComboboxSelected>>', lambda e: self.preview())
        self.map_frame = ttk.Frame(self, padding=10)
        self.map_frame.pack(fill='x')
        for i, (field, (label, _)) in enumerate(FIELDS.items()):
            ttk.Label(self.map_frame, text=label).grid(row=i // 2, column=(i % 2) * 2, sticky='w', padx=4, pady=3)
            box = ttk.Combobox(self.map_frame, state='readonly', width=27)
            box.grid(row=i // 2, column=(i % 2) * 2 + 1, sticky='ew', padx=4)
            self.mapping_boxes[field] = box
        ttk.Label(self, text='供货价与核定售价不可混为一列。到期日可填“长期”；空白为待核对，不默认长期有效。\n多层表头请选包含实际列名的那一行；公式价格请先在Excel计算并保存。', wraplength=900).pack(anchor='w', padx=10)
        self.preview_box = ttk.Frame(self)
        self.preview_box.pack(fill='both', expand=True, padx=10)
        self.info = tk.StringVar()
        ttk.Label(self, textvariable=self.info, wraplength=920).pack(anchor='w', padx=10)
        bottom = ttk.Frame(self, padding=10)
        bottom.pack(fill='x')
        ttk.Button(bottom, text='确认导入本机', command=self.commit).pack(side='right')
        ttk.Button(bottom, text='取消', command=self.destroy).pack(side='right', padx=8)
        self.preview()
        self.transient(parent.winfo_toplevel())
        self.grab_set()

    def preview(self):
        try:
            index = int(self.header.get()) - 1
            rows = self.tables[self.sheet.get()]
            if index < 0 or index >= len(rows):
                raise ValueError('表头行超出范围')
            self.preview_header = index + 1
            headers = rows[index]
            self.choices = ['不导入此字段'] + ['%s · %s' % (i + 1, text(h) or '空表头') for i, h in enumerate(headers)]
            mapping = suggest_mapping(headers)
            for field, box in self.mapping_boxes.items():
                box.configure(values=self.choices)
                box.current(mapping[field] + 1 if field in mapping else 0)
            for widget in self.preview_box.winfo_children():
                widget.destroy()
            tree = table(self.preview_box, self.choices[1:], height=6)
            for row in rows[index + 1:index + 7]:
                tree.insert('', 'end', values=row)
            self.info.set('共 %s 行数据（含可能的空行和说明行）；以下列映射需要你确认。' % (len(rows) - index - 1))
        except Exception as exc:
            self.info.set(str(exc))

    def commit(self):
        try:
            if int(self.header.get()) != self.preview_header:
                raise ValueError('表头行已变更，请先刷新预览')
            mapping = {key: box.current() - 1 for key, box in self.mapping_boxes.items() if box.current() > 0}
            quotes = parse_quotes(self.tables[self.sheet.get()], mapping, self.path, self.sheet.get(), self.preview_header)
            self.callback(quotes)
            self.destroy()
        except Exception as exc:
            messagebox.showerror('无法导入', str(exc), parent=self)


class PricingFrame(ttk.Frame):
    def __init__(self, parent, pool):
        super().__init__(parent, padding=8)
        self.pool = pool
        self.bos = None
        self.data = {'quotes': [], 'bindings': {}, 'catalog': [], 'catalog_time': ''}
        self.source_by_iid, self.candidates_by_iid = {}, {}
        self.results = []
        self.build()

    def build(self):
        actions = ttk.Frame(self)
        actions.pack(fill='x')
        ttk.Button(actions, text='导入采购定价表', command=self.import_file).pack(side='left', padx=3)
        ttk.Button(actions, text='使用已读取商品', command=self.use_current).pack(side='left', padx=3)
        ttk.Button(actions, text='刷新商城及全部门店', command=self.pool.load_all).pack(side='left', padx=3)
        self.info = tk.StringVar(value='先在“微盟商品池”连接商户，再导入采购定价表。所有匹配只保存在本机。')
        ttk.Label(self, textvariable=self.info, wraplength=840).pack(anchor='w', pady=5)
        ttk.Label(self, textvariable=self.pool.status_var, wraplength=840).pack(anchor='w')
        tabs = ttk.Notebook(self)
        tabs.pack(fill='both', expand=True)
        mapping = ttk.Frame(tabs, padding=5)
        tabs.add(mapping, text='采购表与商品匹配')
        row = ttk.Frame(mapping)
        row.pack(fill='x')
        ttk.Label(row, text='采购品名 / 供应商').pack(side='left')
        self.query = tk.StringVar()
        ttk.Entry(row, textvariable=self.query, width=20).pack(side='left', padx=4)
        self.expiry = ttk.Combobox(row, state='readonly', width=15, values=['全部', '有效', '30天内到期', '长期有效', '已过期', '未生效', '未填到期日期', '数据有误'])
        self.expiry.current(0)
        self.expiry.pack(side='left', padx=3)
        self.match_filter = ttk.Combobox(row, state='readonly', width=12, values=['全部', '未匹配', '已确认匹配'])
        self.match_filter.current(0)
        self.match_filter.pack(side='left', padx=3)
        ttk.Button(row, text='筛选', command=self.refresh).pack(side='left')
        self.source = table(mapping, ['采购品名', '品牌', '型号', '单位', '供应商', '采购供货价', '核定售价', '到期日期', '有效期状态', '匹配数', '导入问题'], 8)
        self.source.bind('<<TreeviewSelect>>', lambda e: self.show_candidates())
        ttk.Label(mapping, text='选中上方采购行，再选下方商品并确认。名称相近不自动匹配；一个采购品可确认关联多个门店。').pack(anchor='w')
        row = ttk.Frame(mapping)
        row.pack(fill='x')
        self.candidate_query = tk.StringVar()
        ttk.Entry(row, textvariable=self.candidate_query, width=25).pack(side='left')
        ttk.Button(row, text='搜索候选商品', command=self.show_candidates).pack(side='left', padx=3)
        ttk.Button(row, text='确认关联（仅本机）', command=self.bind_selected).pack(side='left', padx=3)
        ttk.Button(row, text='解除选中采购行的关联', command=self.unbind).pack(side='left', padx=3)
        self.candidate = table(mapping, ['门店', '上架品名', '商品编码', '最低成本', '最高成本', '多规格', '匹配依据'], 6)
        result = ttk.Frame(tabs, padding=5)
        tabs.add(result, text='核对结果与导出')
        row = ttk.Frame(result)
        row.pack(fill='x')
        self.result_filter = ttk.Combobox(row, state='readonly', values=['全部', '成本一致', '成本不一致', '定价有效性待核对', '未匹配', '多规格成本区间待核对', '多个定价冲突', '微盟成本未返回', '商品未在当前快照'], width=23)
        self.result_filter.current(0)
        self.result_filter.pack(side='left')
        ttk.Button(row, text='按当前日期重新核对', command=self.refresh_results).pack(side='left', padx=3)
        ttk.Button(row, text='导出当前核对结果', command=self.export).pack(side='left', padx=3)
        self.result_columns = ['采购品名', '供应商', '门店', '上架品名', '采购供货价', '采购核定售价', '最低成本', '最高成本', '成本差额', '成本核对', '有效期状态', '到期日期', '商城禁售控制', '门店上下架']
        self.result_tree = table(result, self.result_columns, 18)
        ttk.Label(result, text='核定售价用于后续对照真实活动价，不与虚高的商品销售价冒充比较。匹配以本机确认记录为准；多规格没有SKU成本时不判定精确一致。', wraplength=840).pack(anchor='w')

    def select_merchant(self, bos):
        bos = str(bos)
        if self.bos == bos:
            return
        self.bos = bos
        path = self.storage_path()
        if path.exists():
            try:
                self.data = json.loads(path.read_text(encoding='utf-8'))
            except Exception as exc:
                self.bos = None
                raise ValueError('本地定价记录损坏，未覆盖原文件：' + str(exc))
        else:
            self.data = {'quotes': [], 'bindings': {}, 'catalog': [], 'catalog_time': ''}
        self.refresh()

    def storage_path(self):
        if not self.bos or not self.bos.isdigit():
            raise ValueError('请先选择商户')
        return Path(app_dir()) / 'config' / ('pricing_' + self.bos + '.json')

    def save(self):
        path = self.storage_path()
        path.parent.mkdir(exist_ok=True)
        temp = path.with_suffix('.tmp')
        temp.write_text(json.dumps(self.data, ensure_ascii=False, indent=2), encoding='utf-8')
        temp.replace(path)

    def import_file(self):
        if self.pool.busy:
            return
        if not self.bos:
            messagebox.showinfo('先连接商户', '请先在微盟商品池连接并选择商户')
            return
        path = filedialog.askopenfilename(filetypes=[('采购定价表', '*.xlsx *.csv')])
        if path:
            try:
                ImportDialog(self, path, self.accept_quotes)
            except Exception as exc:
                messagebox.showerror('读取失败', str(exc))

    def accept_quotes(self, quotes):
        known = {q['uid'] for q in self.data['quotes']}
        additions = [q for q in quotes if q['uid'] not in known]
        self.data['quotes'].extend(additions)
        self.save()
        self.refresh()
        errors = sum(bool(q.get('errors')) for q in additions)
        self.info.set('新增 %s 行，重复跳过 %s 行，有问题 %s 行（保留原始行，未静默丢弃）。新版本定价不会覆盖旧记录；重叠定价会标记冲突。' % (len(additions), len(quotes) - len(additions), errors))

    def update_catalog(self, records, scope):
        if not self.bos:
            return
        self.data['catalog'] = records
        self.data['catalog_time'] = datetime.now().isoformat(timespec='seconds')
        self.data['catalog_scope'] = scope
        self.save()
        self.refresh()

    def use_current(self):
        if not self.pool.source_records or not self.bos:
            messagebox.showinfo('还没有商品', '请先读取商品池或全部门店')
            return
        records = [(s, r) for s, r in self.pool.source_records if str(s.get('vidType')) == '10']
        if not records:
            messagebox.showinfo('需要门店商品', '请点击“刷新商城及全部门店”，成本需关联门店')
            return
        self.update_catalog(records, '当前已读取门店，可能不是全量')

    def refresh(self):
        self.source.delete(*self.source.get_children())
        self.source_by_iid = {}
        query = self.query.get().strip().casefold()
        for q in self.data['quotes']:
            count = len(self.data['bindings'].get(q['uid'], []))
            if query and query not in (q.get('name', '') + ' ' + q.get('supplier', '')).casefold():
                continue
            if self.expiry.get() not in ('全部', expiry_state(q)):
                continue
            if self.match_filter.get() == '未匹配' and count:
                continue
            if self.match_filter.get() == '已确认匹配' and not count:
                continue
            iid = self.source.insert('', 'end', values=[q.get('name'), q.get('brand'), q.get('model'), q.get('unit'), q.get('supplier'), q.get('cost'), q.get('sale'), q.get('end'), expiry_state(q), count, '；'.join(q.get('errors', []))])
            self.source_by_iid[iid] = q
        self.info.set('采购记录 %s 行；门店商品快照 %s 条；读取时间 %s；范围：%s' % (len(self.data['quotes']), len(self.data['catalog']), self.data.get('catalog_time') or '未读取', self.data.get('catalog_scope', '未读取')))
        self.candidate.delete(*self.candidate.get_children())
        self.refresh_results()

    def selected_quote(self):
        selection = self.source.selection()
        return self.source_by_iid.get(selection[0]) if selection else None

    def show_candidates(self):
        self.candidate.delete(*self.candidate.get_children())
        self.candidates_by_iid = {}
        quote = self.selected_quote()
        if not quote:
            return
        for _, reason, store, row in candidate_products(quote, self.data['catalog'], self.candidate_query.get().strip()):
            price = row.get('goodsPrice') or {}
            iid = self.candidate.insert('', 'end', values=[store.get('vidName'), row.get('title'), row.get('outerGoodsCode') or row.get('goodsCode'), price.get('minCostPrice'), price.get('maxCostPrice'), row.get('isMultiSku'), reason])
            self.candidates_by_iid[iid] = (store, row)

    def bind_selected(self):
        if self.pool.busy:
            return
        quote, selection = self.selected_quote(), self.candidate.selection()
        if not quote or not selection:
            return
        store, row = self.candidates_by_iid[selection[0]]
        prompt = '采购：%s / %s / %s\n型号：%s，单位：%s\n微盟：%s / %s\n\n请核实供应商、规格及计价单位相同。确认只保存本机关联，不改微盟商品。' % (quote.get('supplier'), quote.get('name'), quote.get('spec', ''), quote.get('model', ''), quote.get('unit', ''), store.get('vidName'), row.get('title'))
        if not messagebox.askyesno('确认商品对应关系', prompt):
            return
        key = product_key(store, row)
        bindings = self.data['bindings'].setdefault(quote['uid'], [])
        if key not in bindings:
            bindings.append(key)
        self.save()
        self.refresh()

    def unbind(self):
        if self.pool.busy:
            return
        quote = self.selected_quote()
        if quote and messagebox.askyesno('解除本机关联', '只解除这条采购记录的商品关联；原采购记录和微盟数据不变。'):
            self.data['bindings'].pop(quote['uid'], None)
            self.save()
            self.refresh()

    def refresh_results(self):
        catalog = {product_key(s, r): (s, r) for s, r in self.data['catalog']}
        active = {}
        for q in self.data['quotes']:
            if expiry_state(q) in ('有效', '30天内到期', '长期有效'):
                for key in self.data['bindings'].get(q['uid'], []):
                    active.setdefault(key, []).append(q['uid'])
        results = []
        for q in self.data['quotes']:
            keys = self.data['bindings'].get(q['uid'], [])
            for key in keys or [None]:
                store, row = catalog.get(key, ({}, {}))
                result = compare_quote(q, store, row)
                if key is None:
                    result['成本核对'] = '未匹配'
                elif key not in catalog:
                    result['成本核对'] = '商品未在当前快照'
                elif len(active.get(key, [])) > 1 and q['uid'] in active[key]:
                    result['成本核对'] = '多个定价冲突'
                    result['成本差额'] = None
                result['核对日期'] = date.today().isoformat()
                result['商品快照时间'] = self.data.get('catalog_time', '')
                result['商品快照范围'] = self.data.get('catalog_scope', '')
                if self.result_filter.get() in ('全部', result['成本核对']):
                    results.append(result)
        self.results = results
        self.result_tree.delete(*self.result_tree.get_children())
        for result in results:
            self.result_tree.insert('', 'end', values=[text(result.get(col)) for col in self.result_columns])

    def export(self):
        self.refresh_results()
        if not self.results:
            return
        path = filedialog.asksaveasfilename(defaultextension='.xlsx', initialfile='采购定价核对.xlsx')
        if path:
            try:
                export_comparisons(path, self.results)
                self.info.set('已导出当前核对筛选结果 %s 行：%s' % (len(self.results), path))
            except Exception as exc:
                messagebox.showerror('导出失败', str(exc))


def export_comparisons(path, results):
    # Application runtime export, using the existing desktop dependency.
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    ws = wb.active
    ws.title = '采购定价核对'
    columns = list(results[0]) if results else []
    ws.append(columns)
    numeric = {'采购供货价', '采购核定售价', '最低成本', '最高成本', '成本差额', '活动最低价', '活动最高价', '目标售价', '价差', '核定售价差额', '接口SKU成本', '活动价', '加价率', '毛利率'}
    for record in results:
        ws.append([float(record[col]) if col in numeric and record.get(col) is not None else record.get(col) for col in columns])
    for row in ws:
        for cell in row:
            if isinstance(cell.value, str):
                cell.data_type = 's'
    for i, col in enumerate(columns, 1):
        ws.column_dimensions[get_column_letter(i)].width = 35 if '品名' in col else 22
        ws.cell(1, i).font = Font(bold=True, color='FFFFFF')
        ws.cell(1, i).fill = PatternFill('solid', fgColor='244A70')
        if col in numeric:
            for row in ws.iter_rows(min_row=2, min_col=i, max_col=i):
                row[0].number_format = '0.00%' if col in ('加价率', '毛利率') else '0.00'
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    wb.save(path)
