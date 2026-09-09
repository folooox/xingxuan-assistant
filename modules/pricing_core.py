"""Local purchasing quotations and conservative, reviewable product matching.

No remote writes. Names only suggest candidates, never authorize a match.
"""
import csv
import hashlib
import json
import re
import unicodedata
from datetime import date, datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from difflib import SequenceMatcher
from pathlib import Path

FIELDS = {
    'name': ('采购品名', ['采购品名', '商品名称', '产品名称', '品名', '名称']),
    'supplier': ('供应商', ['供应商', '供应商名称', '供货商', '商家', '中标供应商']),
    'code': ('统一编码（可空）', ['统一编码', '商品编码', '产品编码', '商品条码', '条码']),
    'spec': ('规格 / 单位', ['规格', '规格型号', '商品规格', '规格单位']),
    'model': ('型号', ['型号', '产品型号']),
    'brand': ('品牌', ['品牌']),
    'unit': ('计价单位', ['单位', '计价单位']),
    'cost': ('采购供货价', ['供货价', '采购价', '成本价', '采购成本', '采购供货价', '定价（一件代发含税运）']),
    'sale': ('采购核定售价', ['核定售价', '定价', '零售价', '销售价', '售价', '采购核定售价', '星选销售价']),
    'start': ('生效日期（可空）', ['生效日期', '开始日期', '起始日期']),
    'end': ('到期日期', ['到期日期', '过期时间', '截止日期', '有效期至', '有效期', '到期时间']),
    'owner': ('采购老师', ['采购老师', '采购人', '负责人', '采购员']),
}


def text(value):
    if value is None:
        return ''
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return str(value).strip()


def normalized(value):
    return re.sub(r'[\W_]+', '', unicodedata.normalize('NFKC', text(value)).casefold())


def money(value):
    if value is None or text(value) == '':
        return None
    try:
        amount = Decimal(text(value).replace('￥', '').replace('¥', '').replace(',', ''))
        if not amount.is_finite() or amount < 0:
            raise InvalidOperation
        return amount
    except InvalidOperation:
        raise ValueError('价格必须是不小于0的数字')


def date_value(value):
    raw = text(value)
    if not raw:
        return None
    if raw in ('长期', '长期有效', '永久'):
        return 'long_term'
    raw = raw.replace('/', '-').replace('.', '-').replace('年', '-').replace('月', '-').replace('日', '')
    match = re.fullmatch(r'(\d{4})-(\d{1,2})-(\d{1,2})(?:[ T].*)?', raw)
    if not match:
        raise ValueError('日期需含完整年月日，或填写“长期”')
    try:
        return date(*map(int, match.groups())).isoformat()
    except ValueError:
        raise ValueError('日期无效')


def date_range(value):
    """Actual supplied table uses YYYY.M.D-YYYY.M.D in one 有效期 column."""
    raw = text(value)
    parts = re.findall(r'\d{4}[./年-]\d{1,2}[./月-]\d{1,2}日?', raw)
    if len(parts) == 2:
        remainder = raw.replace(parts[0], '').replace(parts[1], '').strip()
        if remainder not in ('-', '至', '到', '~', '～', '—', '–'):
            raise ValueError('有效期含额外说明，请核对原始记录')
        return date_value(parts[0]), date_value(parts[1])
    return None, date_value(raw)


def expiry_state(quote, today=None):
    today = today or date.today()
    if quote.get('errors'):
        return '数据有误'
    start, end = quote.get('start'), quote.get('end')
    if start and start != 'long_term' and date.fromisoformat(start) > today:
        return '未生效'
    if not end:
        return '未填到期日期'
    if end == 'long_term':
        return '长期有效'
    days = (date.fromisoformat(end) - today).days
    return '已过期' if days < 0 else ('30天内到期' if days <= 30 else '有效')


def load_tables(path):
    """Read values only; never execute macros, formulas or links."""
    path = Path(path)
    if path.suffix.lower() == '.csv':
        for encoding in ('utf-8-sig', 'gb18030'):
            try:
                with path.open(encoding=encoding, newline='') as stream:
                    return {path.stem: list(csv.reader(stream))}
            except UnicodeDecodeError:
                continue
        raise ValueError('无法识别CSV编码')
    if path.suffix.lower() != '.xlsx':
        raise ValueError('请将文件另存为 .xlsx 或 .csv 后导入；不执行宏')
    from openpyxl import load_workbook
    book = load_workbook(path, read_only=False, data_only=True, keep_links=False)
    try:
        tables = {}
        for sheet in book:
            raw = [[text(v) for v in row] for row in sheet.iter_rows(values_only=True)]
            rows = TableRows([list(row) for row in raw])
            rows.original_rows = raw
            rows.merged_rows = {}
            for merged in sheet.merged_cells.ranges:
                value = raw[merged.min_row - 1][merged.min_col - 1]
                for r in range(merged.min_row, merged.max_row + 1):
                    for c in range(merged.min_col, merged.max_col + 1):
                        rows[r - 1][c - 1] = value
                    rows.merged_rows.setdefault(r, []).append(str(merged))
            tables[sheet.title] = rows
        return tables
    finally:
        book.close()


class TableRows(list):
    """Keep original empty merged cells alongside their explicit merged value."""
    pass


def suggest_mapping(headers):
    # Exact header aliases only; ambiguous columns require manual selection.
    result = {}
    for field, (_, aliases) in FIELDS.items():
        candidates = [i for i, header in enumerate(headers) if normalized(header) in {normalized(a) for a in aliases}]
        if len(candidates) == 1:
            result[field] = candidates[0]
    return result


def parse_quotes(rows, mapping, source, sheet, header_row=1):
    if 'name' not in mapping or 'supplier' not in mapping:
        raise ValueError('必须选择采购品名和供应商列')
    if 'cost' not in mapping and 'sale' not in mapping:
        raise ValueError('至少选择采购供货价或采购核定售价之一')
    if len(set(mapping.values())) != len(mapping):
        raise ValueError('同一列不能同时对应多个字段，请分别指定供货价和核定售价')
    if header_row < 1 or header_row > len(rows):
        raise ValueError('表头行超出文件范围')
    quotes = []
    for number, row in enumerate(rows[header_row:], header_row + 1):
        if not any(text(v) for v in row):
            continue
        q = {field: text(row[index]) if index < len(row) else '' for field, index in mapping.items()}
        original = getattr(rows, 'original_rows', rows)[number - 1]
        q.update(source=Path(source).name, sheet=sheet, line=number, raw=[text(v) for v in original], errors=[])
        q['merged_cells'] = sorted(getattr(rows, 'merged_rows', {}).get(number, []))
        if not q.get('name'):
            q['errors'].append('缺少采购品名')
        if not q.get('supplier'):
            q['errors'].append('缺少供应商')
        for field in ('cost', 'sale'):
            try:
                value = money(q.get(field))
                q[field] = str(value) if value is not None else None
            except ValueError as exc:
                q['errors'].append(FIELDS[field][0] + '：' + str(exc))
                q[field] = None
        if q.get('cost') is None and q.get('sale') is None:
            q['errors'].append('两种价格均为空或无效（公式请先在Excel重新计算并保存）')
        try:
            inferred_start, inferred_end = date_range(q.get('end'))
            if inferred_start:
                if q.get('start') and date_value(q['start']) != inferred_start:
                    raise ValueError('生效日期与有效期区间冲突')
                q['start'] = inferred_start
            q['end'] = inferred_end
        except ValueError as exc:
            q['errors'].append('有效期：' + str(exc))
            q['end'] = None
        for field in ('start',):
            try:
                q[field] = date_value(q.get(field))
            except ValueError as exc:
                q['errors'].append(FIELDS[field][0] + '：' + str(exc))
                q[field] = None
        if q.get('start') == 'long_term':
            q['errors'].append('生效日期不能填写长期')
        elif q.get('start') and q.get('end') not in (None, 'long_term') and q['start'] > q['end']:
            q['errors'].append('生效日期晚于到期日期')
        q['uid'] = hashlib.sha256(json.dumps(q, ensure_ascii=False, sort_keys=True).encode()).hexdigest()[:24]
        quotes.append(q)
    if not quotes:
        raise ValueError('没有读到数据行，请检查工作表和表头行')
    return quotes


def product_key(store, row):
    return str(store.get('vid', '')) + ':' + str(row.get('goodsId', ''))


def candidate_products(quote, records, keyword='', limit=50):
    name = normalized(quote.get('name'))
    code = text(quote.get('code'))
    result = []
    for store, row in records:
        actual_name = normalized(row.get('title'))
        actual_code = text(row.get('outerGoodsCode') or row.get('goodsCode'))
        if keyword and normalized(keyword) not in normalized(' '.join([text(row.get('title')), text(store.get('vidName')), actual_code])):
            continue
        score = SequenceMatcher(None, name, actual_name).ratio() if name else 0
        reason = '名称相似，仅作候选'
        if name and (name in actual_name or actual_name in name):
            score = max(score, .75)
        if code and code == actual_code:
            score, reason = 2, '编码一致，仍需确认供应商和规格'
        model = normalized(quote.get('model'))
        if model and len(model) > 2 and model in actual_name:
            score += .6
            reason = '型号一致，需确认供应商与规格'
        if normalized(quote.get('supplier')) == normalized(store.get('vidName')):
            score += .1
        if score > .15 or keyword:
            result.append((score, reason, store, row))
    return sorted(result, key=lambda item: (-item[0], product_key(item[2], item[3])))[:limit]


def price_check(cost_low, cost_high, sale_low, sale_high, markup='1.15'):
    try:
        values = [money(v) for v in (cost_low, cost_high, sale_low, sale_high)]
    except ValueError:
        return {'state': '价格数据异常'}
    lo, hi, sale_lo, sale_hi = values
    if any(v is None for v in values):
        return {'state': '缺少成本或活动价'}
    if lo > hi or sale_lo > sale_hi:
        return {'state': '价格区间异常'}
    if lo == 0:
        return {'state': '零成本待核实'}
    if lo != hi or sale_lo != sale_hi:
        return {'state': '多规格区间待核对'}
    target = (lo * Decimal(markup)).quantize(Decimal('.01'), rounding=ROUND_HALF_UP)
    return {'state': '活动价倒挂' if sale_lo < lo else ('未达加价目标' if sale_lo < target else '达到加价目标'),
            'target': str(target), 'difference': str(sale_lo - lo),
            'markup': str(sale_lo / lo - 1),
            'margin': str((sale_lo - lo) / sale_lo) if sale_lo else None}


def compare_quote(quote, store, row, today=None):
    from modules.product_status import status_values
    status = expiry_state(quote, today)
    price = row.get('goodsPrice') or {}
    lo, hi = price.get('minCostPrice'), price.get('maxCostPrice')
    expected = money(quote.get('cost'))
    difference = None
    if status in ('已过期', '未生效', '未填到期日期', '数据有误'):
        cost_state = '定价有效性待核对'
    elif expected is None:
        cost_state = '未填采购供货价'
    elif lo is None or hi is None:
        cost_state = '微盟成本未返回'
    elif money(lo) != money(hi):
        cost_state = '多规格成本区间待核对'
    else:
        difference = money(lo) - expected
        cost_state = '成本一致' if difference == 0 else '成本不一致'
    return {
        '采购品名': quote.get('name', ''), '供应商': quote.get('supplier', ''),
        '统一编码': quote.get('code', ''), '采购规格': quote.get('spec', ''),
        '采购型号': quote.get('model', ''), '采购品牌': quote.get('brand', ''), '计价单位': quote.get('unit', ''),
        '采购供货价': quote.get('cost'), '采购核定售价': quote.get('sale'),
        '生效日期': quote.get('start'), '到期日期': quote.get('end'), '有效期状态': status,
        '采购老师': quote.get('owner', ''), '门店': store.get('vidName', ''),
        '商品ID': text(row.get('goodsId')), '上架品名': row.get('title', ''),
        '最低成本': lo, '最高成本': hi, '成本差额': str(difference) if difference is not None else None,
        '成本核对': cost_state, '商城禁售控制': status_values(row, store)[0],
        '门店上下架': status_values(row, store)[1],
        '来源文件': quote.get('source'), '工作表': quote.get('sheet'), '来源行': quote.get('line'),
        '导入问题': '；'.join(quote.get('errors', [])),
        '合并单元格来源': '；'.join(quote.get('merged_cells', [])),
    }
