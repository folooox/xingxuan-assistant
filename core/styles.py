# -*- coding: utf-8 -*-
"""Excel 样式公共件：边框、对齐、常用填充色"""
from openpyxl.styles import Alignment, PatternFill, Border, Side, Font

THIN_BORDER = Border(left=Side("thin"), right=Side("thin"),
                     top=Side("thin"), bottom=Side("thin"))
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
FILL_HEADER = PatternFill("solid", fgColor="E0E0E0")
FILL_SUM = PatternFill("solid", fgColor="FFC000")
GROUP_FILLS = [PatternFill("solid", fgColor=c)
               for c in ["E6F3FF", "E6FFF2", "FFF0E6", "F2E6FF", "E6FFFF"]]
BOLD = Font(bold=True)


def apply_border_align(ws, border=THIN_BORDER, align=CENTER):
    """整表加边框；未设置对齐的单元格给居中"""
    for row in ws.iter_rows():
        for cell in row:
            cell.border = border
            if not cell.alignment.horizontal:
                cell.alignment = align
