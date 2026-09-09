import unittest
from modules.weimob_admin_client import WeimobClient, WeimobAPIError
from modules.pool import goods_values

class PoolTests(unittest.TestCase):
    def test_filters(self):
        from decimal import Decimal
        from modules.pool_filter import matches, cost_bound
        row={'title':'商品','isOnline':True,'goodsPrice':{'minCostPrice':2,'maxCostPrice':4}}
        self.assertTrue(matches(row,{'vidName':'门店甲'},{'low':Decimal(3),'store':'甲'}))
        self.assertFalse(matches(row,{}, {'low':Decimal(5)}))
        self.assertFalse(matches({}, {}, {'low':Decimal(0)}))
        self.assertTrue(matches({}, {}, {'cost_state':'成本未返回'}))
        self.assertFalse(matches(row,{}, {'online':'已下架'}))
        with self.assertRaises(ValueError): cost_bound('NaN')

    def test_cost_range_and_zero(self):
        self.assertEqual(goods_values({'goodsPrice':{'minCostPrice':0,'maxCostPrice':0}})[1], '0')
        self.assertEqual(goods_values({'goodsPrice':{'minCostPrice':2,'maxCostPrice':4}})[1], '2 ～ 4')
        self.assertEqual(goods_values({})[1], '未返回')

    def test_export_numeric_cost_and_text_id(self):
        import tempfile
        from pathlib import Path
        from openpyxl import load_workbook
        from modules.pool_export import export_records
        with tempfile.TemporaryDirectory() as tmp:
            path=Path(tmp)/'check.xlsx'
            export_records(path,'test',[({'vidName':'store'},{'goodsId':123456789012345678,'title':'=1+1','goodsPrice':{'minCostPrice':0,'maxCostPrice':2}})])
            wb=load_workbook(path)
            ws=wb['商品门店成本']
            self.assertEqual(ws['H2'].value,0)
            self.assertEqual(ws['I2'].value,2)
            self.assertEqual(ws['D2'].data_type,'s')
            self.assertEqual(ws['E2'].data_type,'s')
            wb.close()

    def test_missing_is_not_zero(self):
        self.assertEqual(goods_values({})[3], '未返回')
        self.assertEqual(goods_values({'goodsStock':{'goodsStockNum':0}})[3], '0')

    def test_exact_id_and_failed_store(self):
        client = WeimobClient('unit-test')
        def products(bos, store, *args):
            if store['vid'] == 2:
                raise WeimobAPIError('读取失败')
            return {'totalCount':2, 'pageList':[{'goodsId':9}, {'goodsId':7}]}
        client.products = products
        result = client.product_stores(1, [{'vid':1,'vidType':10}, {'vid':2,'vidType':10}], {'goodsId':7,'title':'同名'})
        self.assertEqual(result['matches'][0]['product']['goodsId'], 7)
        self.assertEqual(len(result['errors']), 1)

    def test_store_pagination(self):
        client = WeimobClient('unit-test')
        pages = iter([{'subListVos':[{'vid':1}], 'more':True,'scrollId':'a'},
                      {'subListVos':[{'vid':2}], 'more':False}])
        client.call = lambda *args: next(pages)
        self.assertEqual(len(client.stores(1)), 2)

if __name__ == '__main__':
    unittest.main()
