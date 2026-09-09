import json
from pathlib import Path
from modules.weimob_admin_client import WeimobClient
from modules.pool_export import export_records
from openpyxl import load_workbook
c=json.loads(Path('config/weimob_admin.json').read_text(encoding='utf-8'))
client=WeimobClient(c.get('apikey') or c.get('token'))
bos=c['current_merchant']['bos_id']
stores=client.stores(bos)
records,errors=client.all_store_products(bos,stores,lambda s:print(s,flush=True))
out=Path('dist/全部门店商品成本.xlsx')
export_records(out,c['current_merchant']['name'],records,errors,scope='全部可访问门店实测导出')
wb=load_workbook(out,read_only=True)
assert wb['商品门店成本'].max_row==len(records)+1
costs=sum(1 for s,r in records if (r.get('goodsPrice') or {}).get('minCostPrice') is not None)
print(json.dumps({'rows':len(records),'with_cost':costs,'failed_stores':len(errors)},ensure_ascii=False))
wb.close()
